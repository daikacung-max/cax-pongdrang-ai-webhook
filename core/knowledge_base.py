"""Curated, durable knowledge documents for the CAX AI assistant.

New material is always a draft. Only an authenticated officer can approve it,
and only approved chunks are exposed to retrieval. The database stores extracted
text and provenance; source files are never served back to citizens.
"""

from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import unicodedata
import uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from config import DATABASE_URL, DB_PATH
from core.history import _postgres_pool

MAX_TEXT_CHARS = 2_000_000
MAX_UPLOAD_BYTES = 12 * 1024 * 1024
CHUNK_CHARS = 1400
CHUNK_OVERLAP = 180
STATUSES = {"draft", "approved", "archived"}
_SCHEMA_READY = False
_SEARCH_STOPWORDS = {
    "toi", "anh", "chi", "cho", "hoi", "muon", "can", "gi", "nao", "the",
    "la", "va", "co", "khong", "duoc", "huong", "dan", "giup", "ve",
    "vui", "long", "xin", "hay", "tai", "sao", "nhu", "nay", "do",
    "mot", "cac", "nhung", "voi", "tren", "trong", "bao", "nhieu",
}


def _now():
    return datetime.now(timezone.utc).isoformat()


@contextmanager
def _sqlite():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(DB_PATH), timeout=10)
    con.row_factory = sqlite3.Row
    try:
        yield con
        con.commit()
    finally:
        con.close()


def init_schema():
    global _SCHEMA_READY
    if _SCHEMA_READY:
        return
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("""
                    CREATE TABLE IF NOT EXISTS knowledge_documents (
                        id TEXT PRIMARY KEY, title TEXT NOT NULL, source_url TEXT NOT NULL DEFAULT '',
                        issuer TEXT NOT NULL DEFAULT '', document_number TEXT NOT NULL DEFAULT '',
                        source_index INTEGER NOT NULL, effective_from TEXT NOT NULL DEFAULT '',
                        effective_to TEXT NOT NULL DEFAULT '', checked_at TEXT NOT NULL DEFAULT '',
                        status TEXT NOT NULL DEFAULT 'draft', original_name TEXT NOT NULL DEFAULT '',
                        sha256 TEXT NOT NULL, content TEXT NOT NULL, created_at TEXT NOT NULL,
                        updated_at TEXT NOT NULL, approved_at TEXT NOT NULL DEFAULT '',
                        approved_by TEXT NOT NULL DEFAULT '', notes TEXT NOT NULL DEFAULT ''
                    )
                """)
                cur.execute("CREATE INDEX IF NOT EXISTS idx_knowledge_status_source ON knowledge_documents(status, source_index)")
                cur.execute("CREATE TABLE IF NOT EXISTS knowledge_chunks (id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES knowledge_documents(id) ON DELETE CASCADE, chunk_index INTEGER NOT NULL, content TEXT NOT NULL, search_text TEXT NOT NULL)")
                cur.execute("CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_document ON knowledge_chunks(document_id, chunk_index)")
        _SCHEMA_READY = True
        return
    with _sqlite() as con:
        con.execute("""
            CREATE TABLE IF NOT EXISTS knowledge_documents (
                id TEXT PRIMARY KEY, title TEXT NOT NULL, source_url TEXT NOT NULL DEFAULT '',
                issuer TEXT NOT NULL DEFAULT '', document_number TEXT NOT NULL DEFAULT '',
                source_index INTEGER NOT NULL, effective_from TEXT NOT NULL DEFAULT '',
                effective_to TEXT NOT NULL DEFAULT '', checked_at TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT 'draft', original_name TEXT NOT NULL DEFAULT '',
                sha256 TEXT NOT NULL, content TEXT NOT NULL, created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL, approved_at TEXT NOT NULL DEFAULT '',
                approved_by TEXT NOT NULL DEFAULT '', notes TEXT NOT NULL DEFAULT ''
            )
        """)
        con.execute("CREATE INDEX IF NOT EXISTS idx_knowledge_status_source ON knowledge_documents(status, source_index)")
        con.execute("CREATE TABLE IF NOT EXISTS knowledge_chunks (id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES knowledge_documents(id) ON DELETE CASCADE, chunk_index INTEGER NOT NULL, content TEXT NOT NULL, search_text TEXT NOT NULL)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_document ON knowledge_chunks(document_id, chunk_index)")
        con.commit()
    _SCHEMA_READY = True


def _clean_text(text):
    text = str(text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("Tài liệu vượt giới hạn 2 triệu ký tự sau khi trích xuất.")
    if len(text) < 40:
        raise ValueError("Không trích xuất được đủ nội dung để lập chỉ mục.")
    return text


def extract_upload(filename, raw):
    """Extract text from supported, size-limited uploads without saving binaries."""
    name = Path(str(filename or "tai-lieu")).name
    suffix = Path(name).suffix.lower()
    if not raw or len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("Tệp trống hoặc vượt giới hạn 12 MB.")
    if suffix in {".txt", ".md", ".text"}:
        text = raw.decode("utf-8-sig", errors="replace")
    elif suffix == ".pdf":
        from pypdf import PdfReader
        from io import BytesIO
        text = "\n\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(raw)).pages)
    elif suffix == ".docx":
        from io import BytesIO
        from docx import Document
        doc = Document(BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raise ValueError("Chỉ nhận tệp PDF, DOCX, TXT hoặc MD.")
    return name, _clean_text(text), hashlib.sha256(raw).hexdigest()


def _insert_sqlite(values, doc_id, chunks):
    columns = "id,title,source_url,issuer,document_number,source_index,effective_from,effective_to,checked_at,status,original_name,sha256,content,created_at,updated_at,approved_at,approved_by,notes"
    with _sqlite() as con:
        con.execute(f"INSERT INTO knowledge_documents ({columns}) VALUES ({','.join('?' for _ in values)})", values)
        con.executemany("INSERT INTO knowledge_chunks(id,document_id,chunk_index,content,search_text) VALUES (?,?,?,?,?)", [
            (f"{doc_id}:chunk:{i}", doc_id, i, chunk, _search_normalize(str(values[1]) + " " + chunk))
            for i, chunk in enumerate(chunks)
        ])
        con.commit()


def create_document(*, title, content, source_index, source_url="", issuer="", document_number="", effective_from="", effective_to="", checked_at="", original_name="", sha256="", notes=""):
    title = str(title or "").strip()[:300]
    content = _clean_text(content)
    source_index = int(source_index)
    if not title or not 1 <= source_index <= 19:
        raise ValueError("Cần tên tài liệu và nhóm nguồn từ 1 đến 19.")
    for field, value in (("ngày hiệu lực", effective_from), ("ngày hết hiệu lực", effective_to), ("ngày kiểm tra", checked_at)):
        if value and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value)):
            raise ValueError(f"{field.capitalize()} phải theo dạng YYYY-MM-DD.")
    if effective_from and effective_to and effective_to < effective_from:
        raise ValueError("Ngày hết hiệu lực không thể trước ngày bắt đầu hiệu lực.")
    source_url = str(source_url or "").strip()
    if source_url:
        parsed_url = urlparse(source_url)
        if parsed_url.scheme.lower() != "https" or not parsed_url.netloc:
            raise ValueError("Liên kết nguồn phải dùng HTTPS và có tên miền hợp lệ.")
    now = _now()
    doc_id = "KB_" + uuid.uuid4().hex
    digest = sha256 or hashlib.sha256(content.encode("utf-8")).hexdigest()
    chunks = _chunks(content)
    values = (doc_id, title, source_url[:1000], str(issuer or "").strip()[:200],
        str(document_number or "").strip()[:200], source_index, str(effective_from or ""),
        str(effective_to or ""), str(checked_at or ""), "draft", str(original_name or "")[:255],
        digest, content, now, now, "", "", str(notes or "").strip()[:2000])
    init_schema()
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("INSERT INTO knowledge_documents (id,title,source_url,issuer,document_number,source_index,effective_from,effective_to,checked_at,status,original_name,sha256,content,created_at,updated_at,approved_at,approved_by,notes) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", values)
                cur.executemany("INSERT INTO knowledge_chunks(id,document_id,chunk_index,content,search_text) VALUES (%s,%s,%s,%s,%s)", [
                    (f"{doc_id}:chunk:{i}", doc_id, i, chunk, _search_normalize(title + " " + chunk))
                    for i, chunk in enumerate(chunks)
                ])
    else:
        _insert_sqlite(values, doc_id, chunks)
    return get_document(doc_id)


def _row_dict(row, postgres=False, include_content=False):
    if postgres:
        keys = ("id", "title", "source_url", "issuer", "document_number", "source_index", "effective_from", "effective_to", "checked_at", "status", "original_name", "sha256", "content", "created_at", "updated_at", "approved_at", "approved_by", "notes")
        item = dict(zip(keys, row))
    else:
        item = dict(row)
    if not include_content:
        item.pop("content", None)
    return item


def get_document(doc_id, include_content=False):
    init_schema()
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("SELECT id,title,source_url,issuer,document_number,source_index,effective_from,effective_to,checked_at,status,original_name,sha256,content,created_at,updated_at,approved_at,approved_by,notes FROM knowledge_documents WHERE id=%s", (str(doc_id),))
                row = cur.fetchone()
        return _row_dict(row, True, include_content) if row else None
    with _sqlite() as con:
        row = con.execute("SELECT * FROM knowledge_documents WHERE id=?", (str(doc_id),)).fetchone()
    return _row_dict(row, False, include_content) if row else None


def list_documents(status=None, limit=200):
    init_schema()
    limit = min(max(int(limit), 1), 500)
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                if status in STATUSES:
                    cur.execute("SELECT id,title,source_url,issuer,document_number,source_index,effective_from,effective_to,checked_at,status,original_name,sha256,content,created_at,updated_at,approved_at,approved_by,notes FROM knowledge_documents WHERE status=%s ORDER BY updated_at DESC LIMIT %s", (status, limit))
                else:
                    cur.execute("SELECT id,title,source_url,issuer,document_number,source_index,effective_from,effective_to,checked_at,status,original_name,sha256,content,created_at,updated_at,approved_at,approved_by,notes FROM knowledge_documents ORDER BY updated_at DESC LIMIT %s", (limit,))
                rows = cur.fetchall()
        return [_row_dict(row, True) for row in rows]
    with _sqlite() as con:
        if status in STATUSES:
            rows = con.execute("SELECT * FROM knowledge_documents WHERE status=? ORDER BY updated_at DESC LIMIT ?", (status, limit)).fetchall()
        else:
            rows = con.execute("SELECT * FROM knowledge_documents ORDER BY updated_at DESC LIMIT ?", (limit,)).fetchall()
    return [_row_dict(row) for row in rows]


def set_status(doc_id, status, actor=""):
    if status not in STATUSES:
        raise ValueError("Trạng thái không hợp lệ.")
    current = get_document(doc_id)
    if status == "approved":
        if not current:
            return None
        if not current.get("checked_at"):
            raise ValueError("Cần ghi ngày cán bộ đối chiếu trước khi duyệt.")
        if not (current.get("source_url") or current.get("issuer") or current.get("original_name")):
            raise ValueError("Cần ghi cơ quan ban hành, liên kết nguồn hoặc tệp gốc trước khi duyệt.")
    now = _now()
    init_schema()
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("UPDATE knowledge_documents SET status=%s, updated_at=%s, approved_at=CASE WHEN %s='approved' THEN %s ELSE approved_at END, approved_by=CASE WHEN %s='approved' THEN %s ELSE approved_by END WHERE id=%s", (status, now, status, now, status, str(actor or "")[:120], str(doc_id)))
                found = cur.rowcount
    else:
        with _sqlite() as con:
            cur = con.execute("UPDATE knowledge_documents SET status=?, updated_at=?, approved_at=CASE WHEN ?='approved' THEN ? ELSE approved_at END, approved_by=CASE WHEN ?='approved' THEN ? ELSE approved_by END WHERE id=?", (status, now, status, now, status, str(actor or "")[:120], str(doc_id)))
            found = cur.rowcount
            con.commit()
    return get_document(doc_id) if found else None


def _chunks(text):
    paragraphs = [x.strip() for x in re.split(r"\n\s*\n", text) if x.strip()]
    chunks, current = [], ""
    for paragraph in paragraphs:
        while len(paragraph) > CHUNK_CHARS:
            cut = paragraph.rfind(" ", 0, CHUNK_CHARS)
            cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
            piece, paragraph = paragraph[:cut].strip(), paragraph[cut:].strip()
            if current:
                chunks.append(current)
                current = ""
            chunks.append(piece)
        candidate = paragraph if not current else current + "\n\n" + paragraph
        if len(candidate) > CHUNK_CHARS and current:
            chunks.append(current)
            current = (current[-CHUNK_OVERLAP:] + "\n\n" + paragraph).strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def _search_normalize(value):
    value = unicodedata.normalize("NFD", str(value or "").lower())
    value = "".join(char for char in value if unicodedata.category(char) != "Mn")
    return value.replace("đ", "d")


def search_approved(query, source_index, limit=4):
    """Small, deterministic lexical search bounded to one approved Notebook source."""
    query = str(query or "").strip()
    normalized_query = _search_normalize(query)
    tokens = [token for token in dict.fromkeys(re.findall(r"\w{2,}", normalized_query)) if token not in _SEARCH_STOPWORDS][:12]
    if not tokens:
        return []
    init_schema()
    patterns = ["%" + token + "%" for token in tokens]
    if DATABASE_URL:
        with _postgres_pool().connection() as con:
            with con.cursor() as cur:
                where = " OR ".join("c.search_text LIKE %s" for _ in patterns)
                cur.execute(f"SELECT d.id,d.title,d.source_url,d.issuer,d.document_number,d.effective_from,d.effective_to,d.checked_at,c.chunk_index,c.content,c.search_text FROM knowledge_documents d JOIN knowledge_chunks c ON c.document_id=d.id WHERE d.status='approved' AND d.source_index=%s AND (d.effective_from='' OR d.effective_from<=%s) AND (d.effective_to='' OR d.effective_to>=%s) AND ({where}) ORDER BY d.updated_at DESC LIMIT 800", (int(source_index), datetime.now(timezone.utc).date().isoformat(), datetime.now(timezone.utc).date().isoformat(), *patterns))
                rows = cur.fetchall()
    else:
        with _sqlite() as con:
            where = " OR ".join("c.search_text LIKE ?" for _ in patterns)
            as_of = datetime.now(timezone.utc).date().isoformat()
            rows = con.execute(f"SELECT d.id,d.title,d.source_url,d.issuer,d.document_number,d.effective_from,d.effective_to,d.checked_at,c.chunk_index,c.content,c.search_text FROM knowledge_documents d JOIN knowledge_chunks c ON c.document_id=d.id WHERE d.status='approved' AND d.source_index=? AND (d.effective_from='' OR d.effective_from<=?) AND (d.effective_to='' OR d.effective_to>=?) AND ({where}) ORDER BY d.updated_at DESC LIMIT 800", (int(source_index), as_of, as_of, *patterns)).fetchall()
    as_of = datetime.now(timezone.utc).date().isoformat()
    scored = []
    for row in rows:
        doc_id, title, source_url, issuer, number, start, end, checked, i, chunk, blob = tuple(row)
        matched = sum(1 for token in tokens if token in blob)
        if matched >= min(2, len(tokens)):
            scored.append((matched, doc_id, i, title, source_url, issuer, number, start, end, checked, chunk))
    scored.sort(key=lambda x: (x[0], x[1]), reverse=True)
    result = []
    for score, doc_id, index, title, url, issuer, number, start, end, checked, chunk in scored[:min(max(int(limit), 1), 8)]:
        result.append({
            "id": f"{doc_id}:chunk:{index}", "document_id": doc_id,
            "unit_type": "knowledge_chunk", "title": title, "text": chunk,
            "document_title": title, "document_number": number,
            "document_issuer": issuer, "source_path": url,
            "effective_from": start, "effective_to": end,
            "_knowledge_score": score,
        })
    return result


def document_for_unit(unit_id):
    doc_id = str(unit_id or "").split(":chunk:", 1)[0]
    if not doc_id.startswith("KB_"):
        return None
    return get_document(doc_id)
