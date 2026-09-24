"""Citizen-requested document drafting and printable DOCX export.

The module is intentionally opt-in. Ordinary legal/TTHC questions continue
through AI Core unchanged; this path activates only when a citizen explicitly
asks to print, fill, draft or export a form/document.

A form is a temporary task inside a wider conversation, not a sticky chatbot
mode. Short structured field replies may continue the form, while a clear new
topic immediately returns control to AI Core.
"""

from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from io import BytesIO
from urllib.parse import quote, unquote

from cryptography.fernet import Fernet, InvalidToken
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from config import DATABASE_URL, HISTORY_HMAC_SECRET, UNIT_NAME
from core import history


PUBLIC_BASE_URL = os.getenv(
    "PUBLIC_BASE_URL", "https://cax-pongdrang-ai.onrender.com"
).rstrip("/")
FORM_LINK_TTL_SECONDS = max(300, int(os.getenv("FORM_LINK_TTL_SECONDS", "1800")))


def _norm(value: str) -> str:
    value = str(value or "").strip().casefold()
    table = str.maketrans(
        "àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ",
        "aaaaaaaaaaaaaaaaaeeeeeeeeeeeiiiiiooooooooooooooooouuuuuuuuuuuyyyyyd",
    )
    return " ".join(value.translate(table).split())


def _fernet() -> Fernet:
    secret = (
        os.getenv("FORM_LINK_SECRET", "").strip()
        or str(HISTORY_HMAC_SECRET or "").strip()
        # Render services provision DATABASE_URL for the durable history store.
        # Use it only as a last-resort instance-bound key so a missing optional
        # form secret cannot silently bypass the form engine. Never fall back to
        # a public or source-controlled constant.
        or os.getenv("DATABASE_URL", "").strip()
        # Older Render services may not yet have the Blueprint-managed history
        # variables. Provider keys are already protected runtime secrets, so
        # they keep the opt-in form path functional until the dedicated secret
        # is provisioned; rotating the provider key invalidates old form links.
        or os.getenv("GROQ_API_KEY", "").strip()
        or os.getenv("OPENAI_API_KEY", "").strip()
    )
    if not secret:
        raise RuntimeError("form_link_secret_missing")
    digest = hashlib.sha256(("cax-form-link:" + secret).encode("utf-8")).digest()
    return Fernet(base64.urlsafe_b64encode(digest))


def encode_payload(payload: dict) -> str:
    body = dict(payload or {})
    body["issued_at"] = int(datetime.now(timezone.utc).timestamp())
    raw = json.dumps(body, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return _fernet().encrypt(raw).decode("ascii")


def decode_payload(token: str) -> dict:
    try:
        # Tokens are URL-quoted in download links. Flask decodes path segments
        # before route handling, but callers/tests may pass the literal encoded
        # segment. Accept both representations without weakening Fernet checks.
        decoded_token = unquote(str(token or "").strip())
        raw = _fernet().decrypt(decoded_token.encode("ascii"), ttl=FORM_LINK_TTL_SECONDS)
        return json.loads(raw.decode("utf-8"))
    except (InvalidToken, ValueError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValueError("invalid_or_expired_form_link") from exc


_SHORT_LINK_PREFIX = "f_"


def _store_short_payload(user_id: str, payload: dict) -> str:
    """Store an encrypted form payload behind a compact Zalo-safe URL token."""
    history.init_schema()
    token = _SHORT_LINK_PREFIX + secrets.token_urlsafe(12)
    encrypted_payload = encode_payload(payload)
    key = history.conversation_key(user_id)
    expires_at = history.utc_now() + timedelta(seconds=FORM_LINK_TTL_SECONDS)

    if DATABASE_URL:
        with history._postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("DELETE FROM form_downloads WHERE expires_at < %s", (history.utc_now(),))
                cur.execute(
                    "INSERT INTO form_downloads(token, user_key, payload_token, expires_at) VALUES (%s, %s, %s, %s)",
                    (token, key, encrypted_payload, expires_at),
                )
        return token

    with history._sqlite() as con:
        now_text = history.utc_now().isoformat()
        con.execute("DELETE FROM form_downloads WHERE expires_at < ?", (now_text,))
        con.execute(
            "INSERT INTO form_downloads(token, user_key, payload_token, expires_at) VALUES (?, ?, ?, ?)",
            (token, key, encrypted_payload, expires_at.isoformat()),
        )
    return token


def decode_download_token(token: str) -> dict:
    """Resolve compact server-side links while accepting issued legacy links."""
    token = unquote(str(token or "").strip())
    if not token.startswith(_SHORT_LINK_PREFIX):
        return decode_payload(token)

    history.init_schema()
    if DATABASE_URL:
        with history._postgres_pool().connection() as con:
            with con.cursor() as cur:
                cur.execute("SELECT payload_token, expires_at FROM form_downloads WHERE token=%s", (token,))
                row = cur.fetchone()
                if not row or row[1] < history.utc_now():
                    cur.execute("DELETE FROM form_downloads WHERE token=%s", (token,))
                    raise ValueError("invalid_or_expired_form_link")
                encrypted_payload = row[0]
    else:
        with history._sqlite() as con:
            row = con.execute("SELECT payload_token, expires_at FROM form_downloads WHERE token=?", (token,)).fetchone()
            if not row or str(row["expires_at"]) < history.utc_now().isoformat():
                con.execute("DELETE FROM form_downloads WHERE token=?", (token,))
                raise ValueError("invalid_or_expired_form_link")
            encrypted_payload = row["payload_token"]
    return decode_payload(encrypted_payload)


def _field(lines: list[str], aliases: tuple[str, ...]) -> str:
    normalized_aliases = {_norm(x) for x in aliases}
    for line in lines:
        if ":" not in line:
            continue
        label, value = line.split(":", 1)
        if _norm(label) in normalized_aliases and value.strip():
            return value.strip()
    return ""


def _collect_user_text(history, question: str) -> str:
    chunks = [str(item.get("content") or "") for item in (history or []) if item.get("role") == "user"]
    chunks.append(str(question or ""))
    return "\n".join(chunks)


def _latest_assistant_item(history):
    for item in reversed(history or []):
        if item.get("role") == "assistant":
            return item
    return {}


def _active_form(history) -> str:
    latest = _latest_assistant_item(history)
    meta = latest.get("meta") or {}
    if meta.get("path") == "citizen_form_assistant" and not meta.get("form_ready"):
        form_type = str(meta.get("form_type") or "").strip()
        if form_type in ("ct01", "report"):
            return form_type

    n = _norm(latest.get("content") or "")
    if "vui long gui cac dong sau" in n and "ct01" in n:
        return "ct01"
    if "vui long gui cac dong sau" in n and "don trinh bao" in n:
        return "report"
    if "con thieu" in n and "ct01" in n:
        return "ct01"
    if "con thieu" in n and "don trinh bao" in n:
        return "report"
    return ""


_FORM_FIELD_LABELS = {
    "ho ten", "ho va ten", "ho chu dem va ten", "ngay sinh", "ngay thang nam sinh",
    "gioi tinh", "cccd", "so cccd", "so dinh danh ca nhan", "ddcn", "so dien thoai",
    "dien thoai", "sdt", "noi thuong tru", "thuong tru", "noi tam tru", "tam tru",
    "noi o hien tai", "cho o hien tai", "dia chi hien tai", "nghe nghiep",
    "nghe nghiep noi lam viec", "chu ho", "ho ten chu ho", "quan he voi chu ho",
    "noi dung de nghi", "de nghi", "noi dung dang ky", "dia chi", "noi o", "noi cu tru",
    "thoi gian xay ra", "thoi gian", "ngay gio xay ra", "dia diem xay ra", "dia diem",
    "noi dung su viec", "su viec", "noi dung trinh bao", "tai lieu chung cu", "chung cu",
    "tai lieu kem theo", "yeu cau",
}


def _looks_like_form_field_reply(question: str) -> bool:
    for line in str(question or "").splitlines():
        if ":" not in line:
            continue
        label, value = line.split(":", 1)
        if _norm(label) in _FORM_FIELD_LABELS and value.strip():
            return True
    return False


def _is_form_capability_question(question: str) -> bool:
    n = _norm(question)
    asks_file = any(x in n for x in (
        "xuat file", "tao file", "tai file", "file word", "file world", "file docx", "co file khong",
        "co xuat duoc file", "xuat duoc khong", "in duoc khong",
    ))
    return asks_file and any(x in n for x in ("co", "duoc", "khong", "the nao", "lam sao"))


def _explicit_form_type(question: str) -> str:
    n = _norm(question)
    explicit_action = any(x in n for x in (
        "mau", "bieu mau", "in giup", "in cho", "dien mau", "xuat file",
        "tai file", "soan don", "don trinh bao", "viet don", "lap don", "lam don", "tao don",
    ))
    if not explicit_action:
        return ""
    if any(x in n for x in (
        "don trinh bao", "don to giac", "soan don trinh bao", "viet don trinh bao",
        "lap don trinh bao", "lam don trinh bao", "tao don trinh bao",
    )):
        return "report"
    # A citizen often describes the incident and asks simply "giúp tôi làm
    # đơn".  Treat that as a report only when the same turn contains a clear
    # incident cue; ordinary civil requests must remain on the normal AI path.
    report_cues = (
        "bi trom", "mat tai san", "bi lua", "lua dao", "bi danh", "hanh hung",
        "bi de doa", "danh bac", "ca do", "ma tuy", "gay roi", "pha hoai",
    )
    if any(x in n for x in ("soan don", "viet don", "lap don", "lam don", "tao don")) and any(
        cue in n for cue in report_cues
    ):
        return "report"
    if any(x in n for x in (
        "ct01", "cu tru", "thuong tru", "tam tru", "thay doi thong tin cu tru",
        "dang ky cu tru",
    )):
        return "ct01"
    return ""


def _continue_active_form(question: str, active: str) -> bool:
    n = _norm(question)
    if any(x in n for x in ("huy bieu mau", "huy don", "khong lam nua", "dung lai")):
        return True
    if _looks_like_form_field_reply(question):
        return True
    if _is_form_capability_question(question):
        return True
    if any(x in n for x in (
        "tiep tuc don", "tiep tuc mau", "don nay", "mau nay", "bieu mau nay",
        "dien tiep", "lam tiep", "xuat luon", "tao luon file",
    )):
        return True
    return False


def detect_form_type(question: str, history=None) -> str:
    explicit = _explicit_form_type(question)
    if explicit:
        return explicit

    active = _active_form(history or [])
    if active and _continue_active_form(question, active):
        return active
    return ""


def _ct01_fields(text: str) -> dict:
    lines = [x.strip() for x in str(text or "").splitlines() if x.strip()]
    return {
        "full_name": _field(lines, ("Họ tên", "Họ và tên", "Họ, chữ đệm và tên")),
        "dob": _field(lines, ("Ngày sinh", "Ngày tháng năm sinh")),
        "gender": _field(lines, ("Giới tính",)),
        "personal_id": _field(lines, ("CCCD", "Số CCCD", "Số định danh cá nhân", "ĐDCN")),
        "phone": _field(lines, ("Số điện thoại", "Điện thoại", "SĐT")),
        "permanent_address": _field(lines, ("Nơi thường trú", "Thường trú")),
        "temporary_address": _field(lines, ("Nơi tạm trú", "Tạm trú")),
        "current_address": _field(lines, ("Nơi ở hiện tại", "Chỗ ở hiện tại", "Địa chỉ hiện tại")),
        "occupation": _field(lines, ("Nghề nghiệp", "Nghề nghiệp, nơi làm việc")),
        "owner_name": _field(lines, ("Chủ hộ", "Họ tên chủ hộ")),
        "owner_relation": _field(lines, ("Quan hệ với chủ hộ",)),
        "request_content": _field(lines, ("Nội dung đề nghị", "Đề nghị", "Nội dung đăng ký")),
    }


def _natural_value(text: str, patterns: tuple[str, ...]) -> str:
    """Return an explicitly volunteered value from conversational prose.

    This is deliberately conservative: it recognizes a value only after a
    citizen-facing cue (for example, ``tôi tên là``).  It never asks a model to
    invent missing identity information.
    """
    for pattern in patterns:
        match = re.search(pattern, str(text or ""), flags=re.IGNORECASE | re.MULTILINE)
        if match:
            value = re.sub(r"\s+", " ", match.group(1)).strip(" ,.;:-")
            if value:
                return value
    return ""


def _natural_incident_content(text: str) -> str:
    """Keep the citizen's own incident narrative when it is in free prose."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in str(text or "").splitlines()]
    cues = (
        "bi trom", "mat", "bi lua", "lua dao", "bi danh", "hanh hung", "de doa",
        "danh bac", "ca do", "ma tuy", "gay roi", "pha hoai", "chuyen khoan",
    )
    for line in reversed(lines):
        normalized = _norm(line)
        if not normalized:
            continue
        if ":" in line:
            label = line.split(":", 1)[0]
            if _norm(label) in _FORM_FIELD_LABELS:
                continue
        if any(cue in normalized for cue in cues):
            # Remove only the drafting request, preserving the factual words
            # supplied after it in the same natural-language sentence.
            value = re.sub(
                r"^\s*(?:toi\s+(?:muon|can)\s+)?(?:giup\s+toi\s+)?(?:soan|viet|lap|lam|tao)\s+don\s+(?:trinh\s+bao|to\s+giac)\s*(?:cho\s+toi)?\s*[:,.-]?\s*",
                "",
                line,
                flags=re.IGNORECASE,
            ).strip()
            return value or line
    return ""


def _report_fields(text: str) -> dict:
    lines = [x.strip() for x in str(text or "").splitlines() if x.strip()]
    fields = {
        "full_name": _field(lines, ("Họ tên", "Họ và tên")),
        "dob": _field(lines, ("Ngày sinh", "Ngày tháng năm sinh")),
        "personal_id": _field(lines, ("CCCD", "Số CCCD", "Số định danh cá nhân")),
        "address": _field(lines, ("Địa chỉ", "Nơi ở", "Nơi cư trú")),
        "phone": _field(lines, ("Số điện thoại", "Điện thoại", "SĐT")),
        "incident_time": _field(lines, ("Thời gian xảy ra", "Thời gian", "Ngày giờ xảy ra")),
        "incident_place": _field(lines, ("Địa điểm xảy ra", "Địa điểm")),
        "incident_content": _field(lines, ("Nội dung sự việc", "Sự việc", "Nội dung trình báo")),
        "evidence": _field(lines, ("Tài liệu chứng cứ", "Chứng cứ", "Tài liệu kèm theo")),
        "request_content": _field(lines, ("Đề nghị", "Yêu cầu", "Nội dung đề nghị")),
    }
    # Permit a natural conversation such as "Tôi tên là ...; hôm qua tôi bị
    # lừa chuyển khoản ...".  Labelled fields above take precedence.
    fields["full_name"] = fields["full_name"] or _natural_value(
        text, (r"\b(?:tôi\s+tên(?:\s+là)?|tên\s+tôi\s+là|họ\s+tên\s+tôi\s+là)\s+([^,.;\n]{2,80})",)
    )
    fields["personal_id"] = fields["personal_id"] or _natural_value(
        text, (r"\b(?:cccd|căn\s+cước|số\s+định\s+danh(?:\s+cá\s+nhân)?)\s*(?:là)?\s*[:#-]?\s*(\d{9,12})",)
    )
    fields["phone"] = fields["phone"] or _natural_value(
        text, (r"\b(?:số\s+điện\s+thoại|điện\s+thoại|sđt)\s*(?:là)?\s*[:#-]?\s*((?:\+84|0)\d{9,10})",)
    )
    fields["address"] = fields["address"] or _natural_value(
        text, (r"\b(?:địa\s+chỉ(?:\s+của\s+tôi)?|tôi\s+(?:đang\s+)?ở)\s*(?:là|tại)?\s*[:,-]?\s*([^\n.;]{4,160})",)
    )
    fields["incident_content"] = fields["incident_content"] or _natural_incident_content(text)
    return fields


def _missing(fields: dict, required: tuple[str, ...]) -> list[str]:
    return [name for name in required if not str(fields.get(name) or "").strip()]


CT01_LABELS = {
    "full_name": "Họ tên",
    "dob": "Ngày sinh",
    "personal_id": "Số định danh cá nhân/CCCD",
    "current_address": "Nơi ở hiện tại",
    "request_content": "Nội dung đề nghị",
}
REPORT_LABELS = {
    "full_name": "Họ tên",
    "personal_id": "Số định danh cá nhân/CCCD",
    "address": "Địa chỉ",
    "incident_content": "Nội dung sự việc",
    "request_content": "Đề nghị",
}
REPORT_DRAFT_REQUIRED = ("incident_content",)


def _blank_requested(question: str) -> bool:
    n = _norm(question)
    return any(x in n for x in ("mau trong", "ban trong", "in mau", "cho mau", "tai mau")) and not any(
        x in n for x in ("dien", "ghi thong tin", "soan theo thong tin")
    )


def _download_url(form_type: str, fields: dict, user_id: str) -> str:
    token = _store_short_payload(user_id, {"form_type": form_type, "fields": fields})
    filename = "ct01" if form_type == "ct01" else "don-trinh-bao"
    return f"{PUBLIC_BASE_URL}/forms/download/{quote(token, safe='')}/{filename}.docx"


def handle_form_request(user_id: str, question: str, history=None):
    form_type = detect_form_type(question, history=history)
    if not form_type:
        return None

    n = _norm(question)
    if any(x in n for x in ("huy bieu mau", "huy don", "khong lam nua", "dung lai")):
        return {
            "answer": "Đã dừng phần hỗ trợ biểu mẫu. Anh/chị có thể chuyển sang hỏi nội dung khác bất kỳ lúc nào.",
            "form_type": form_type,
            "ready": False,
        }

    combined = _collect_user_text(history or [], question)
    if form_type == "ct01":
        fields = _ct01_fields(combined)
        if _is_form_capability_question(question):
            missing = _missing(fields, tuple(CT01_LABELS))
            if missing:
                labels = ", ".join(CT01_LABELS[k] for k in missing)
                return {
                    "answer": (
                        "Có. Khi đủ thông tin, tôi có thể xuất file Word CT01 để anh/chị tải về, kiểm tra và in. "
                        "Hiện còn thiếu: " + labels + ". Anh/chị có thể gửi từng mục hoặc gửi nhiều dòng cùng lúc."
                    ),
                    "form_type": "ct01",
                    "ready": False,
                    "missing": missing,
                }
        if _blank_requested(question):
            url = _download_url("ct01", {}, user_id)
            return {
                "answer": (
                    "Tôi đã tạo bản CT01 trống để anh/chị tải về, in và tự điền: " + url +
                    "\nBiểu mẫu cần được kiểm tra lại thông tin trước khi ký/nộp."
                ),
                "form_type": "ct01",
                "ready": True,
                "download_url": url,
            }
        required = tuple(CT01_LABELS)
        missing = _missing(fields, required)
        if missing:
            labels = "\n".join(f"- {CT01_LABELS[k]}:" for k in missing)
            return {
                "answer": (
                    "Tôi đang hỗ trợ điền CT01 (Tờ khai thay đổi thông tin cư trú). "
                    "Tôi không tự điền thay các thông tin anh/chị chưa cung cấp.\n"
                    "Anh/chị có thể gửi một hoặc nhiều mục sau:\n" + labels
                ),
                "form_type": "ct01",
                "ready": False,
                "missing": missing,
            }
        url = _download_url("ct01", fields, user_id)
        return {
            "answer": (
                "Đã đủ dữ liệu tối thiểu để lập bản CT01 hỗ trợ điền. Anh/chị tải file Word tại: " + url +
                "\nTrước khi in, ký hoặc nộp, vui lòng kiểm tra lại toàn bộ thông tin và bổ sung các mục còn để trống nếu thuộc trường hợp của mình."
            ),
            "form_type": "ct01",
            "ready": True,
            "download_url": url,
        }

    fields = _report_fields(combined)
    if _is_form_capability_question(question) and not fields.get("incident_content"):
        missing = _missing(fields, REPORT_DRAFT_REQUIRED)
        if missing:
            labels = ", ".join(REPORT_LABELS[k] for k in missing)
            return {
                "answer": (
                    "Có. Tôi có thể tạo file Word Đơn trình báo từ nội dung anh/chị kể. "
                    "Hiện còn thiếu: " + labels + ". Anh/chị có thể kể tự nhiên hoặc gửi từng mục/một lúc nhiều dòng."
                ),
                "form_type": "report",
                "ready": False,
                "missing": missing,
            }

    # A Word draft can be useful before the citizen has every identity/detail
    # field.  Require only their own account of the incident; the remaining
    # fields stay blank in the document rather than being guessed by AI.
    missing = _missing(fields, REPORT_DRAFT_REQUIRED)
    if missing:
        labels = "\n".join(f"- {REPORT_LABELS[k]}:" for k in missing)
        return {
            "answer": (
                "Tôi đang hỗ trợ soạn Đơn trình báo gửi " + UNIT_NAME + ". "
                "Tôi chỉ dùng dữ liệu anh/chị cung cấp, không tự suy đoán diễn biến vụ việc.\n"
                "Anh/chị hãy kể ngắn gọn sự việc đã xảy ra (ai, việc gì, khi nào/nơi nào nếu nhớ). "
                "Có thể nhắn tự nhiên, không cần theo mẫu, hoặc gửi:\n" + labels +
                "\nNếu có, có thể bổ sung: Họ tên, CCCD, địa chỉ, số điện thoại, thời gian, địa điểm, tài liệu chứng cứ và đề nghị của anh/chị. "
                "Nếu muốn chuyển sang nội dung khác, anh/chị cứ hỏi bình thường; hệ thống sẽ tự rời chế độ soạn đơn."
            ),
            "form_type": "report",
            "ready": False,
            "missing": missing,
        }
    url = _download_url("report", fields, user_id)
    return {
        "answer": (
            "Tôi đã soạn bản Đơn trình báo Word từ đúng các thông tin anh/chị cung cấp. Tải file tại: " + url +
            "\nCác mục anh/chị chưa cung cấp được để trống để tự điền, không phải thông tin do AI suy đoán. "
            "Anh/chị cần đọc lại nội dung, sửa chi tiết nếu cần và ký xác nhận trước khi sử dụng."
        ),
        "form_type": "report",
        "ready": True,
        "download_url": url,
    }


def _set_default_font(document: Document):
    styles = document.styles
    style = styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)


def _title(document: Document, text: str):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)


def _line(document: Document, label: str, value: str = ""):
    p = document.add_paragraph()
    r = p.add_run(label + ": ")
    r.bold = True
    p.add_run(str(value or ""))


def _make_ct01(fields: dict) -> bytes:
    d = Document()
    _set_default_font(d)
    _title(d, "TỜ KHAI THAY ĐỔI THÔNG TIN CƯ TRÚ (CT01)")
    p = d.add_paragraph("Bản hỗ trợ điền thông tin để người dân kiểm tra, in và ký")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph("Kính gửi: Cơ quan đăng ký cư trú có thẩm quyền")
    _line(d, "Họ, chữ đệm và tên", fields.get("full_name", ""))
    _line(d, "Ngày, tháng, năm sinh", fields.get("dob", ""))
    _line(d, "Giới tính", fields.get("gender", ""))
    _line(d, "Số định danh cá nhân/CCCD", fields.get("personal_id", ""))
    _line(d, "Số điện thoại liên hệ", fields.get("phone", ""))
    _line(d, "Nơi thường trú", fields.get("permanent_address", ""))
    _line(d, "Nơi tạm trú", fields.get("temporary_address", ""))
    _line(d, "Nơi ở hiện tại", fields.get("current_address", ""))
    _line(d, "Nghề nghiệp, nơi làm việc", fields.get("occupation", ""))
    _line(d, "Họ tên chủ hộ", fields.get("owner_name", ""))
    _line(d, "Quan hệ với chủ hộ", fields.get("owner_relation", ""))
    _line(d, "Nội dung đề nghị", fields.get("request_content", ""))
    d.add_paragraph("Tôi cam đoan những thông tin kê khai trên là đúng sự thật và chịu trách nhiệm về nội dung đã kê khai.")
    p = d.add_paragraph("Người kê khai\n(Ký, ghi rõ họ tên)")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph(
        "Lưu ý: File này do hệ thống hỗ trợ điền theo thông tin người dân cung cấp. "
        "Cần đối chiếu biểu mẫu/cách thức tiếp nhận hiện hành và kiểm tra lại trước khi ký, nộp."
    )
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def _make_report(fields: dict) -> bytes:
    d = Document()
    _set_default_font(d)
    _title(d, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    _title(d, "Độc lập - Tự do - Hạnh phúc")
    _title(d, "ĐƠN TRÌNH BÁO")
    d.add_paragraph(f"Kính gửi: {UNIT_NAME}")
    _line(d, "Tôi tên", fields.get("full_name", ""))
    _line(d, "Ngày sinh", fields.get("dob", ""))
    _line(d, "Số định danh cá nhân/CCCD", fields.get("personal_id", ""))
    _line(d, "Địa chỉ", fields.get("address", ""))
    _line(d, "Số điện thoại", fields.get("phone", ""))
    d.add_paragraph("Nay tôi làm đơn này trình báo sự việc như sau:")
    _line(d, "Thời gian xảy ra", fields.get("incident_time", ""))
    _line(d, "Địa điểm xảy ra", fields.get("incident_place", ""))
    _line(d, "Nội dung sự việc", fields.get("incident_content", ""))
    _line(d, "Tài liệu, chứng cứ kèm theo", fields.get("evidence", ""))
    _line(d, "Đề nghị", fields.get("request_content", ""))
    d.add_paragraph("Tôi cam đoan nội dung trình báo trên là đúng theo những gì tôi biết và chịu trách nhiệm về nội dung đã trình bày.")
    p = d.add_paragraph("Người trình báo\n(Ký, ghi rõ họ tên)")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph(
        "Lưu ý: Đây là bản dự thảo do hệ thống hỗ trợ soạn từ dữ liệu người dân cung cấp; "
        "người trình báo phải đọc lại, chỉnh sửa nếu cần và ký xác nhận trước khi sử dụng."
    )
    out = BytesIO()
    d.save(out)
    return out.getvalue()


def render_docx(payload: dict) -> tuple[bytes, str]:
    form_type = str((payload or {}).get("form_type") or "")
    fields = (payload or {}).get("fields") or {}
    if form_type == "ct01":
        return _make_ct01(fields), "CT01-ho-tro-dien.docx"
    if form_type == "report":
        return _make_report(fields), "Don-trinh-bao.docx"
    raise ValueError("unsupported_form_type")
